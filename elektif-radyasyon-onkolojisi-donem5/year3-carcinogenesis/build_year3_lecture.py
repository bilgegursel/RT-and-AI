#!/usr/bin/env python3
"""Year-3 medical lecture: Chemical carcinogens, tobacco, radiation — visual PPTX + Word."""

from pathlib import Path
import shutil
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from lxml import etree
from docx import Document
from docx.shared import Pt as DocPt, RGBColor as DocRGB, Cm, Inches as DocIn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn as doc_qn

OUT = Path(__file__).resolve().parent
ASSETS = OUT / "assets"
DESKTOP_EDU = Path("/home/ubuntu/Desktop/Eğitim")
ARTIFACTS = Path("/opt/cursor/artifacts")

NAVY = RGBColor(0x14, 0x32, 0x42)
INK = RGBColor(0x1F, 0x2A, 0x33)
TEAL = RGBColor(0x1F, 0x7A, 0x6E)
CORAL = RGBColor(0xD4, 0x5D, 0x3A)
AMBER = RGBColor(0xE0, 0xA1, 0x00)
SLATE = RGBColor(0x5A, 0x6B, 0x76)
FOG = RGBColor(0xF3, 0xF6, 0xF7)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LINE = RGBColor(0xD0, 0xD9, 0xDE)
SOFT_TEAL = RGBColor(0xE6, 0xF3, 0xF1)
SOFT_CORAL = RGBColor(0xFB, 0xEE, 0xE9)
SOFT_AMBER = RGBColor(0xFF, 0xF4, 0xD6)
BLUE = RGBColor(0x2E, 0x75, 0xB6)

# Collect notes for Word export
LECTURE_NOTES = []  # list of (slide_no, title, notes)


def font(run, size=18, bold=False, color=INK):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = "Calibri"


def bg(slide, color=FOG):
    f = slide.background.fill
    f.solid()
    f.fore_color.rgb = color


def shp(slide, kind, l, t, w, h, fill, line=None):
    s = slide.shapes.add_shape(kind, l, t, w, h)
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
    return s


def rect(slide, l, t, w, h, fill, line=None):
    return shp(slide, MSO_SHAPE.RECTANGLE, l, t, w, h, fill, line)


def rrect(slide, l, t, w, h, fill, line=None):
    return shp(slide, MSO_SHAPE.ROUNDED_RECTANGLE, l, t, w, h, fill, line)


def oval(slide, l, t, w, h, fill, line=None):
    return shp(slide, MSO_SHAPE.OVAL, l, t, w, h, fill, line)


def chevron(slide, l, t, w, h, fill):
    return shp(slide, MSO_SHAPE.CHEVRON, l, t, w, h, fill)


def rarrow(slide, l, t, w, h, fill):
    return shp(slide, MSO_SHAPE.RIGHT_ARROW, l, t, w, h, fill)


def darrow(slide, l, t, w, h, fill):
    return shp(slide, MSO_SHAPE.DOWN_ARROW, l, t, w, h, fill)


def tb(slide, l, t, w, h, text, size=18, bold=False, color=INK, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(l, t, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    font(r, size, bold, color)
    return box


def shape_text(shape, text, size=14, bold=True, color=WHITE):
    tf = shape.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    if p.runs:
        p.runs[0].text = text
        font(p.runs[0], size, bold, color)
    else:
        r = p.add_run()
        r.text = text
        font(r, size, bold, color)


def bullets(slide, l, t, w, h, items, size=16, color=INK, spacing=7):
    box = slide.shapes.add_textbox(l, t, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(spacing)
        r = p.add_run()
        r.text = "•  " + item
        font(r, size, False, color)
    return box


def footer(slide, section, n, total):
    tb(slide, Inches(0.4), Inches(7.15), Inches(10.5), Inches(0.28),
       f"Year 3 Medicine  ·  Carcinogenesis Lecture  ·  {section}", 11, False, SLATE)
    tb(slide, Inches(11.5), Inches(7.15), Inches(1.5), Inches(0.28),
       f"{n}/{total}", 11, False, SLATE, PP_ALIGN.RIGHT)


def title_bar(slide, title):
    rect(slide, Inches(0), Inches(0), Inches(13.333), Inches(0.14), TEAL)
    tb(slide, Inches(0.5), Inches(0.3), Inches(12.3), Inches(0.5), title, 24, True, NAVY)
    rect(slide, Inches(0.5), Inches(0.88), Inches(1.3), Inches(0.05), CORAL)


def add_notes(slide, title, notes, n):
    notes_slide = slide.notes_slide
    tf = notes_slide.notes_text_frame
    tf.text = notes
    LECTURE_NOTES.append((n, title, notes))


def add_fade_transition(slide):
    """Add a simple fade transition to make the deck feel more dynamic."""
    nsmap = {"p": "http://schemas.openxmlformats.org/presentationml/2006/main"}
    sld = slide._element
    # Remove existing transition
    for child in list(sld):
        if child.tag.endswith("transition"):
            sld.remove(child)
    tr = etree.SubElement(sld, qn("p:transition"))
    tr.set("spd", "med")
    etree.SubElement(tr, qn("p:fade"))


def new_prs():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    return prs


def blank(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def pic(slide, name, l, t, w, h):
    path = ASSETS / name
    if path.exists():
        slide.shapes.add_picture(str(path), l, t, w, h)
        return True
    return False


# ---------- slide constructors ----------

def make_title(prs, n, total):
    s = blank(prs)
    bg(s, NAVY)
    oval(s, Inches(-1.2), Inches(-1), Inches(4), Inches(4), RGBColor(0x1A, 0x45, 0x58))
    oval(s, Inches(10.8), Inches(4.2), Inches(3.8), Inches(3.8), RGBColor(0x1A, 0x45, 0x58))
    rect(s, Inches(0), Inches(5.35), Inches(13.333), Inches(2.15), TEAL)
    for i, (lab, col) in enumerate([("CHEMICALS", CORAL), ("TOBACCO", AMBER), ("RADIATION", BLUE)]):
        ch = rrect(s, Inches(0.7 + i * 2.7), Inches(4.5), Inches(2.4), Inches(0.5), col)
        shape_text(ch, lab, 12, True, WHITE)
    tb(s, Inches(0.7), Inches(1.7), Inches(12), Inches(1.6),
       "Carcinogenesis:\nChemicals, Tobacco & Radiation", 36, True, WHITE)
    tb(s, Inches(0.7), Inches(3.6), Inches(12), Inches(0.5),
       "Year 3 Medical Students  ·  Mechanisms, Exposures, Prevention", 18, False, RGBColor(0xC8, 0xE6, 0xE1))
    tb(s, Inches(0.7), Inches(5.7), Inches(12), Inches(1.2),
       "Prof. Dr. Ş. Bilge Gürsel\nOndokuz Mayıs University Faculty of Medicine  ·  Department of Radiation Oncology",
       15, False, WHITE)
    add_notes(s, "Title",
              "Welcome students. Introduce yourself and the three pillars of today’s lecture: "
              "chemical carcinogens, tobacco (including other tobacco products and sidestream smoke), "
              "and radiation carcinogenesis. Emphasize that the same biological story — DNA injury, "
              "clonal selection, and progression — links all three. Tell them speaker notes are also "
              "provided as a Word document for study.", n)
    add_fade_transition(s)
    footer(s, "Introduction", n, total)
    return s


def make_objectives(prs, n, total):
    s = blank(prs)
    bg(s, FOG)
    title_bar(s, "Learning Objectives")
    objs = [
        ("01", "Explain genotoxic vs non-genotoxic chemical carcinogenesis"),
        ("02", "Classify carcinogens (IARC) and list major environmental/occupational agents"),
        ("03", "Describe exposure routes and practical protection strategies"),
        ("04", "Detail tobacco smoke constituents, products, MSS vs SSS, and mechanisms"),
        ("05", "Outline radiation types, DNA injury pathways, environmental & medical exposure"),
    ]
    for i, (num, text) in enumerate(objs):
        y = 1.25 + i * 1.05
        rrect(s, Inches(0.6), Inches(y), Inches(12.1), Inches(0.9), WHITE, LINE)
        oval(s, Inches(0.85), Inches(y + 0.15), Inches(0.6), Inches(0.6), TEAL if i % 2 == 0 else CORAL)
        tb(s, Inches(0.85), Inches(y + 0.28), Inches(0.6), Inches(0.4), num, 12, True, WHITE, PP_ALIGN.CENTER)
        tb(s, Inches(1.7), Inches(y + 0.25), Inches(10.6), Inches(0.5), text, 16, False, INK)
    add_notes(s, "Learning Objectives",
              "Read the five objectives aloud. Tell students that by the end they should be able to "
              "explain mechanisms, classify exposures, and counsel patients about tobacco and radiation risks. "
              "This is a Year-3 foundational lecture — focus on mechanisms and clinical relevance, not memorizing every chemical.", n)
    add_fade_transition(s)
    footer(s, "Introduction", n, total)
    return s


def make_content(prs, title, items, section, n, total, notes, tip=None):
    s = blank(prs)
    bg(s, FOG)
    title_bar(s, title)
    rrect(s, Inches(0.5), Inches(1.2), Inches(12.3), Inches(5.5 if tip else 5.6), WHITE, LINE)
    rect(s, Inches(0.5), Inches(1.2), Inches(0.14), Inches(5.5 if tip else 5.6), TEAL)
    bullets(s, Inches(0.9), Inches(1.5), Inches(11.5), Inches(5.0), items, 17)
    if tip:
        rrect(s, Inches(0.5), Inches(6.4), Inches(12.3), Inches(0.5), SOFT_CORAL, LINE)
        tb(s, Inches(0.7), Inches(6.48), Inches(11.9), Inches(0.35), "★  " + tip, 12, True, CORAL)
    add_notes(s, title, notes, n)
    add_fade_transition(s)
    footer(s, section, n, total)
    return s


def make_section(prs, num, title, subtitle, n, total, notes):
    s = blank(prs)
    bg(s, TEAL)
    oval(s, Inches(10.2), Inches(-1), Inches(4.5), Inches(4.5), RGBColor(0x28, 0x8A, 0x7C))
    chip = rrect(s, Inches(0.8), Inches(2.3), Inches(2.4), Inches(0.55), WHITE)
    shape_text(chip, f"PART {num}", 14, True, TEAL)
    tb(s, Inches(0.8), Inches(3.1), Inches(11.5), Inches(1.2), title, 34, True, WHITE)
    tb(s, Inches(0.8), Inches(4.5), Inches(11.5), Inches(0.8), subtitle, 18, False, RGBColor(0xD8, 0xF0, 0xEC))
    add_notes(s, title, notes, n)
    add_fade_transition(s)
    footer(s, title, n, total)
    return s


def make_image_slide(prs, title, image, caption, section, n, total, notes):
    s = blank(prs)
    bg(s, FOG)
    title_bar(s, title)
    # image frame
    rrect(s, Inches(0.7), Inches(1.2), Inches(11.9), Inches(5.2), WHITE, LINE)
    ok = pic(s, image, Inches(0.95), Inches(1.4), Inches(11.4), Inches(4.5))
    if not ok:
        tb(s, Inches(1), Inches(3), Inches(11), Inches(1), "[Diagram]", 20, True, SLATE, PP_ALIGN.CENTER)
    tb(s, Inches(0.9), Inches(6.05), Inches(11.5), Inches(0.4), caption, 13, False, SLATE, PP_ALIGN.CENTER)
    add_notes(s, title, notes, n)
    add_fade_transition(s)
    footer(s, section, n, total)
    return s


def make_twocol(prs, title, lh, left, rh, right, section, n, total, notes):
    s = blank(prs)
    bg(s, FOG)
    title_bar(s, title)
    for x, h, items, col in [(0.45, lh, left, TEAL), (6.9, rh, right, CORAL)]:
        rrect(s, Inches(x), Inches(1.2), Inches(5.95), Inches(5.5), WHITE, LINE)
        rect(s, Inches(x), Inches(1.2), Inches(5.95), Inches(0.65), col)
        tb(s, Inches(x + 0.25), Inches(1.32), Inches(5.45), Inches(0.45), h, 17, True, WHITE)
        bullets(s, Inches(x + 0.3), Inches(2.1), Inches(5.4), Inches(4.3), items, 15)
    add_notes(s, title, notes, n)
    add_fade_transition(s)
    footer(s, section, n, total)
    return s


def make_iarc(prs, n, total):
    s = blank(prs)
    bg(s, FOG)
    title_bar(s, "IARC Classification of Carcinogenic Agents")
    levels = [
        (3.6, "Group 1 — Carcinogenic to humans", "Sufficient evidence in humans", TEAL),
        (4.8, "Group 2A — Probably carcinogenic", "Limited human + sufficient animal evidence", CORAL),
        (6.0, "Group 2B — Possibly carcinogenic", "Limited evidence in humans", AMBER),
        (7.2, "Group 3 — Not classifiable", "Inadequate evidence currently", SLATE),
    ]
    for i, (w, title, sub, col) in enumerate(levels):
        x = (13.333 - w) / 2
        y = 1.35 + i * 1.25
        box = rrect(s, Inches(x), Inches(y), Inches(w), Inches(1.05), col)
        tb(s, Inches(x + 0.2), Inches(y + 0.15), Inches(w - 0.4), Inches(0.4), title, 15, True, WHITE, PP_ALIGN.CENTER)
        tb(s, Inches(x + 0.2), Inches(y + 0.55), Inches(w - 0.4), Inches(0.35), sub, 12, False, WHITE, PP_ALIGN.CENTER)
    tb(s, Inches(0.5), Inches(6.55), Inches(12.3), Inches(0.35),
       "Source: IARC Monographs — https://monographs.iarc.who.int/", 11, False, SLATE)
    add_notes(s, "IARC Classification",
              "Walk down the pyramid. Group 1 means established human carcinogen — examples: tobacco smoke, "
              "benzene, asbestos, ionizing radiation, ethanol, aflatoxin. Stress that classification is about "
              "hazard evidence, not individual risk prediction. Always ask: who is exposed, how much, how long?", n)
    add_fade_transition(s)
    footer(s, "Chemical Carcinogens", n, total)
    return s


def make_ipp_chevrons(prs, n, total):
    s = blank(prs)
    bg(s, FOG)
    title_bar(s, "Initiation → Promotion → Progression")
    cols = [
        (0.4, "1. INITIATION", TEAL, ["Irreversible DNA change", "Oncogene / suppressor hits", "DNA-repair failure", "Alone rarely = tumor"]),
        (4.7, "2. PROMOTION", AMBER, ["Clonal expansion", "Mitogenic, reversible", "Needs repeated exposure", "Not mutagenic alone"]),
        (9.0, "3. PROGRESSION", CORAL, ["Genetic instability", "Invasion / metastasis", "Inflammation, angiogenesis", "Clinically malignant"]),
    ]
    for x, title, col, pts in cols:
        ch = chevron(s, Inches(x), Inches(1.3), Inches(3.9), Inches(0.8), col)
        shape_text(ch, title, 14, True, WHITE)
        rrect(s, Inches(x), Inches(2.35), Inches(3.9), Inches(4.3), WHITE, LINE)
        bullets(s, Inches(x + 0.3), Inches(2.7), Inches(3.4), Inches(3.6), pts, 15)
    add_notes(s, "Initiation–Promotion–Progression",
              "This is the master flowchart of the lecture. Initiation is the heritable DNA hit. Promotion "
              "expands that clone under continuous stimulation and is reversible if the promoter stops. "
              "Progression is irreversible evolution toward invasion and metastasis. Classic teaching model: "
              "mouse skin — benzo[a]pyrene initiator then phorbol ester promoter.", n)
    add_fade_transition(s)
    footer(s, "Chemical Carcinogens", n, total)
    return s


def make_entry_cards(prs, n, total):
    s = blank(prs)
    bg(s, FOG)
    title_bar(s, "Types of Exposure to Carcinogens")
    routes = [
        ("INHALATION", "Gases, fumes, dusts, smoke", "Most common occupational route", TEAL, "↑"),
        ("SKIN CONTACT", "Solvents, oils, pesticides", "Absorption through intact skin", CORAL, "◉"),
        ("INGESTION", "Food, water, hand-to-mouth", "Contaminated diet or accidents", AMBER, "↓"),
        ("INJECTION /\nOTHER", "Medical, parenteral, rare", "Iatrogenic or intentional routes", BLUE, "+"),
    ]
    for i, (h, a, b, col, sym) in enumerate(routes):
        x = 0.4 + i * 3.2
        rrect(s, Inches(x), Inches(1.35), Inches(3.0), Inches(5.3), WHITE, LINE)
        oval(s, Inches(x + 0.85), Inches(1.7), Inches(1.3), Inches(1.3), col)
        tb(s, Inches(x + 0.85), Inches(2.05), Inches(1.3), Inches(0.7), sym, 26, True, WHITE, PP_ALIGN.CENTER)
        tb(s, Inches(x + 0.15), Inches(3.3), Inches(2.7), Inches(0.9), h, 15, True, NAVY, PP_ALIGN.CENTER)
        tb(s, Inches(x + 0.2), Inches(4.3), Inches(2.6), Inches(0.9), a, 13, False, SLATE, PP_ALIGN.CENTER)
        rrect(s, Inches(x + 0.25), Inches(5.5), Inches(2.5), Inches(0.85), SOFT_TEAL if i % 2 == 0 else SOFT_CORAL)
        tb(s, Inches(x + 0.35), Inches(5.65), Inches(2.3), Inches(0.6), b, 11, True, INK, PP_ALIGN.CENTER)
    add_notes(s, "Exposure Types",
              "Risk ≈ intensity × duration × susceptibility × route. Inhalation dominates workplaces. "
              "Skin absorption is underestimated for solvents. Ingestion matters for foodborne toxins "
              "like aflatoxin and contaminated water (arsenic). Mention that the same agent can affect "
              "different organs depending on route and metabolism.", n)
    add_fade_transition(s)
    footer(s, "Chemical Carcinogens", n, total)
    return s


def make_protection(prs, n, total):
    s = blank(prs)
    bg(s, FOG)
    title_bar(s, "Protection Against Chemical Carcinogens")
    hierarchy = [
        ("1. Eliminate / Substitute", "Remove the agent or replace with safer alternative", TEAL),
        ("2. Engineering controls", "Ventilation, closed systems, process enclosure", CORAL),
        ("3. Administrative controls", "Training, rotation, exposure limits, labeling", AMBER),
        ("4. PPE (last line)", "Respirators, gloves, goggles, protective clothing", BLUE),
        ("5. Medical surveillance", "Biological monitoring, early detection where indicated", NAVY),
    ]
    for i, (h, b, col) in enumerate(hierarchy):
        y = 1.25 + i * 1.05
        w = 11.5 - i * 0.35
        x = 0.9 + i * 0.15
        box = rrect(s, Inches(x), Inches(y), Inches(w), Inches(0.9), col)
        tb(s, Inches(x + 0.3), Inches(y + 0.12), Inches(w - 0.5), Inches(0.35), h, 15, True, WHITE)
        tb(s, Inches(x + 0.3), Inches(y + 0.48), Inches(w - 0.5), Inches(0.35), b, 12, False, WHITE)
    add_notes(s, "Protection Hierarchy",
              "Teach the industrial-hygiene hierarchy: PPE is not the first answer. Substitution and "
              "engineering controls prevent more cancer than gloves alone. Link to regulation (exposure "
              "limits), workplace rights, and the physician’s role in occupational history-taking.", n)
    add_fade_transition(s)
    footer(s, "Chemical Carcinogens", n, total)
    return s


def make_tobacco_products(prs, n, total):
    s = blank(prs)
    bg(s, FOG)
    title_bar(s, "Tobacco & Nicotine Products — Broader Than Cigarettes")
    cards = [
        ("Cigarettes", "Combusted; highest multi-site cancer evidence", CORAL),
        ("Cigars / pipes", "Oral, laryngeal, esophageal, lung risks", TEAL),
        ("Waterpipe /\nshisha", "Long sessions; high CO & toxicant load", AMBER),
        ("Smokeless\ntobacco", "Oral / pancreatic cancer; no combustion but still carcinogens", BLUE),
        ("Bidís / kreteks", "Often high exposure; regional products", NAVY),
        ("E-cigarettes /\nheated products", "Not risk-free; evolving evidence; not for never-users", SLATE),
    ]
    for i, (h, b, col) in enumerate(cards):
        r, c = divmod(i, 3)
        x = 0.45 + c * 4.2
        y = 1.3 + r * 2.8
        rrect(s, Inches(x), Inches(y), Inches(4.0), Inches(2.55), WHITE, LINE)
        rect(s, Inches(x), Inches(y), Inches(4.0), Inches(0.7), col)
        tb(s, Inches(x + 0.2), Inches(y + 0.15), Inches(3.6), Inches(0.5), h.replace("\n", " "), 15, True, WHITE)
        tb(s, Inches(x + 0.25), Inches(y + 1.0), Inches(3.5), Inches(1.3), b, 14, False, INK)
    add_notes(s, "Other Tobacco Products",
              "Do not let students think ‘only cigarettes cause cancer.’ Combusted products share many "
              "toxicants. Smokeless tobacco still delivers nitrosamines. Waterpipe sessions can deliver "
              "large toxicant doses. E-cigarettes: stress they are not harmless, evidence is evolving, "
              "and they should not be started by never-smokers; cessation context is different.", n)
    add_fade_transition(s)
    footer(s, "Tobacco", n, total)
    return s


def make_mss_sss_cards(prs, n, total):
    s = blank(prs)
    bg(s, FOG)
    title_bar(s, "Mainstream vs Sidestream Smoke")
    # two columns
    rrect(s, Inches(0.45), Inches(1.25), Inches(6.0), Inches(5.5), WHITE, LINE)
    rect(s, Inches(0.45), Inches(1.25), Inches(6.0), Inches(0.75), TEAL)
    tb(s, Inches(0.7), Inches(1.4), Inches(5.5), Inches(0.5), "MAINSTREAM SMOKE (MSS)", 16, True, WHITE)
    bullets(s, Inches(0.75), Inches(2.3), Inches(5.4), Inches(4.1), [
        "Drawn through the cigarette into the smoker’s mouth",
        "Filtered partly by the tobacco rod / filter",
        "Primary exposure for the active smoker",
        "Delivers nicotine + full carcinogen load",
        "Dose often summarized as pack-years",
    ], 14)

    rrect(s, Inches(6.85), Inches(1.25), Inches(6.0), Inches(5.5), WHITE, LINE)
    rect(s, Inches(6.85), Inches(1.25), Inches(6.0), Inches(0.75), CORAL)
    tb(s, Inches(7.1), Inches(1.4), Inches(5.5), Inches(0.5), "SIDESTREAM SMOKE (SSS)", 16, True, WHITE)
    bullets(s, Inches(7.15), Inches(2.3), Inches(5.4), Inches(4.1), [
        "Rises from the smoldering tip between puffs",
        "Generated at lower temperature → different chemistry",
        "Often richer in some toxicants per mg",
        "Major component of second-hand (environmental) smoke",
        "Causes disease in nonsmokers — not ‘mild’ exposure",
    ], 14)
    add_notes(s, "MSS vs SSS",
              "Define clearly. Sidestream smoke is a major part of second-hand smoke. Because combustion "
              "conditions differ, SSS can be disproportionately toxic for some compounds. Emphasize "
              "smoke-free policies protect workers and children. Second-hand smoke is an IARC Group 1 carcinogen.", n)
    add_fade_transition(s)
    footer(s, "Tobacco", n, total)
    return s


def make_radiation_types(prs, n, total):
    s = blank(prs)
    bg(s, FOG)
    title_bar(s, "Types of Radiation Relevant to Carcinogenesis")
    cards = [
        ("Ionizing\nElectromagnetic", "X-rays, γ-rays", "Medical imaging, therapy, background, accidents", TEAL),
        ("Particulate\nIonizing", "α, β, neutrons, protons", "Radon daughters, nuclear, some therapy beams", CORAL),
        ("Non-ionizing\nUV", "UVA, UVB, UVC", "Sunlight, tanning beds → skin cancers / melanoma", AMBER),
        ("Other\nnon-ionizing", "RF, ELF, etc.", "Generally not established like ionizing/UV for solid evidence of cancer", BLUE),
    ]
    for i, (h, ex, note, col) in enumerate(cards):
        x = 0.4 + i * 3.2
        rrect(s, Inches(x), Inches(1.4), Inches(3.05), Inches(5.2), WHITE, LINE)
        rect(s, Inches(x), Inches(1.4), Inches(3.05), Inches(1.5), col)
        tb(s, Inches(x + 0.15), Inches(1.65), Inches(2.75), Inches(1.1), h, 15, True, WHITE, PP_ALIGN.CENTER)
        tb(s, Inches(x + 0.2), Inches(3.2), Inches(2.65), Inches(0.9), ex, 14, True, NAVY, PP_ALIGN.CENTER)
        tb(s, Inches(x + 0.2), Inches(4.3), Inches(2.65), Inches(1.9), note, 13, False, SLATE, PP_ALIGN.CENTER)
    add_notes(s, "Radiation Types",
              "For Year 3, focus on ionizing radiation and UV as established carcinogens. "
              "Particulate radiation (alpha from radon) is highly damaging locally. "
              "Avoid spending lecture time on controversial RF debates — state that evidence is not comparable "
              "to tobacco, asbestos, or ionizing radiation.", n)
    add_fade_transition(s)
    footer(s, "Radiation", n, total)
    return s


def make_env_med_radiation(prs, n, total):
    s = blank(prs)
    bg(s, FOG)
    title_bar(s, "Environmental vs Medical Radiation Exposure")
    rrect(s, Inches(0.45), Inches(1.25), Inches(6.0), Inches(5.5), WHITE, LINE)
    rect(s, Inches(0.45), Inches(1.25), Inches(6.0), Inches(0.7), TEAL)
    tb(s, Inches(0.7), Inches(1.4), Inches(5.5), Inches(0.45), "ENVIRONMENTAL / BACKGROUND", 15, True, WHITE)
    bullets(s, Inches(0.75), Inches(2.2), Inches(5.4), Inches(4.2), [
        "Cosmic rays (higher at altitude / flights)",
        "Terrestrial radionuclides in soil/rock",
        "Radon gas indoors — major natural lung carcinogen",
        "Internal radionuclides (K-40, etc.)",
        "Fallout / accidental releases (rare but historic lessons)",
    ], 14)

    rrect(s, Inches(6.85), Inches(1.25), Inches(6.0), Inches(5.5), WHITE, LINE)
    rect(s, Inches(6.85), Inches(1.25), Inches(6.0), Inches(0.7), CORAL)
    tb(s, Inches(7.1), Inches(1.4), Inches(5.5), Inches(0.45), "MEDICAL EXPOSURE", 15, True, WHITE)
    bullets(s, Inches(7.15), Inches(2.2), Inches(5.4), Inches(4.2), [
        "Diagnostic: X-ray, CT, fluoroscopy, nuclear medicine",
        "Interventional radiology — can deliver high skin/organ doses",
        "Radiotherapy — high therapeutic doses; secondary cancer risk small but real",
        "Principles: justification + optimization (ALARA)",
        "Never withhold indicated imaging from fear alone",
    ], 14)
    add_notes(s, "Environmental vs Medical Radiation",
              "Radon is the teaching pearl for environmental exposure — second leading cause of lung cancer "
              "after smoking in many countries, synergistic with tobacco. For medical exposure, teach "
              "justification and ALARA. CT is useful but not casual. Radiotherapy’s benefit outweighs "
              "secondary cancer risk when indicated; survivors still need long-term awareness.", n)
    add_fade_transition(s)
    footer(s, "Radiation", n, total)
    return s


def make_takeaways(prs, n, total):
    s = blank(prs)
    bg(s, FOG)
    title_bar(s, "Take-Home Messages")
    msgs = [
        ("1", "Carcinogens act by damaging DNA or promoting clonal growth — often both."),
        ("2", "IARC Group 1 agents are established human carcinogens; exposure control saves lives."),
        ("3", "Tobacco is a multi-carcinogen mixture; MSS and SSS both matter; other products are not safe."),
        ("4", "Ionizing radiation and UV are proven carcinogens; risk depends on dose, age, and tissue."),
        ("5", "Prevention hierarchy: eliminate/substitute → engineering → admin → PPE; plus tobacco cessation & ALARA."),
        ("6", "Take an exposure history: occupation, smoking, sun, radon, prior radiation."),
    ]
    for i, (num, msg) in enumerate(msgs):
        y = 1.25 + i * 0.9
        rrect(s, Inches(0.5), Inches(y), Inches(12.3), Inches(0.8), WHITE, LINE)
        oval(s, Inches(0.7), Inches(y + 0.12), Inches(0.55), Inches(0.55), TEAL if i % 2 == 0 else CORAL)
        tb(s, Inches(0.7), Inches(y + 0.22), Inches(0.55), Inches(0.4), num, 14, True, WHITE, PP_ALIGN.CENTER)
        tb(s, Inches(1.5), Inches(y + 0.2), Inches(11.0), Inches(0.45), msg, 14, False, INK)
    add_notes(s, "Take-Home Messages",
              "Close by repeating these six points. Invite questions from each section. Remind students "
              "the Word document contains the full speaking script for revision.", n)
    add_fade_transition(s)
    footer(s, "Summary", n, total)
    return s


def build_pptx():
    prs = new_prs()
    builders = []

    def add(fn):
        builders.append(fn)

    # INTRO
    add(lambda i, t: make_title(prs, i, t))
    add(lambda i, t: make_objectives(prs, i, t))
    add(lambda i, t: make_content(prs, "Why This Matters for Year-3 Doctors", [
        "Most cancers are linked to modifiable environmental and lifestyle exposures",
        "Physicians must recognize occupational and lifestyle carcinogens in the history",
        "Prevention (tobacco, workplace, sun, radon, rational imaging) is clinical work",
        "Mechanistic literacy helps you understand pathology, oncology, and public health",
        "Same DNA-injury biology underlies chemicals, smoke, and radiation",
    ], "Introduction", i, t,
        "Open with curiosity: if inheritance explains <15% of cancers, what explains the rest?",
        tip="Environment + lifestyle dominate population cancer burden."))

    add(lambda i, t: make_content(prs, "Big Picture Numbers", [
        "Lifetime cancer risk in developed countries ≈ 30%",
        "~85–95% of cancers relate to environmental / lifestyle factors",
        "Inherited cancer syndromes: <15%",
        "Avoidable contributors (classic estimates): tobacco ~30%, diet, infections, radiation/UV, occupation, pollution",
        "Geographic differences largely track exposures, not only genes",
    ], "Introduction", i, t,
        "Cite Doll-type avoidable-risk teaching figures as order-of-magnitude guidance, not exact local statistics."))

    # CHEMICALS
    add(lambda i, t: make_section(prs, "01", "Chemical Carcinogens",
                                  "Mechanisms · classification · exposures · protection", i, t,
                                  "We now enter chemical carcinogenesis — the core mechanism block of the lecture."))

    add(lambda i, t: make_content(prs, "What Is a Chemical Carcinogen?", [
        "An agent that increases cancer incidence after exposure (human and/or experimental evidence)",
        "Historic landmark: Percival Pott (1775) — scrotal cancer in chimney sweeps (soot / PAHs)",
        "Modern framework: IARC Monographs integrate human, animal, and mechanistic data",
        "Latency may be years to decades after exposure begins",
        "Dose, route, duration, and host factors determine individual risk",
    ], "Chemical Carcinogens", i, t,
        "Tell the Pott story briefly — it invented environmental cancer thinking."))

    add(lambda i, t: make_iarc(prs, i, t))

    add(lambda i, t: make_image_slide(prs, "Genotoxic vs Non-Genotoxic — Visual Overview",
        "genotoxic_nongenotoxic.jpg",
        "Two mechanistic families — many real-world agents combine features of both",
        "Chemical Carcinogens", i, t,
        "Use this figure as the fork in the road. Genotoxic agents mutate DNA. Non-genotoxic agents "
        "drive proliferation, receptor signaling, cytotoxicity with regeneration, or oxidative stress. "
        "Ask a student for one example of each before you reveal lists."))

    add(lambda i, t: make_twocol(prs, "Mechanisms in Detail",
        "Genotoxic",
        ["DNA-reactive parent compound or metabolite",
         "DNA adducts, strand breaks, chromosomal injury",
         "Often mutagenic in short assays",
         "Direct-acting (alkylators) vs procarcinogens (need CYP450 activation)",
         "Examples: PAHs, nitrosamines, aflatoxin, benzene metabolites"],
        "Non-genotoxic",
        ["Do not directly mutate DNA as first step",
         "Mitogens / cytotoxicants / receptor agonists",
         "Promote clonal expansion of initiated cells",
         "Harder to detect — often need rodent bioassays",
         "Examples: some hormones, certain promoters, chronic irritation pathways"],
        "Chemical Carcinogens", i, t,
        "Stress metabolic activation: many chemicals are safe until CYP450 creates the ultimate carcinogen. "
        "Inter-individual enzyme differences partly explain susceptibility."))

    add(lambda i, t: make_image_slide(prs, "Initiation–Promotion–Progression — Diagram",
        "carcinogenesis_ipp_diagram.jpg",
        "The master sequence of chemical (and many physical) carcinogenesis",
        "Chemical Carcinogens", i, t,
        "Walk left to right on the figure. Initiation = irreversible DNA hit. Promotion = expansion. "
        "Progression = invasion/metastasis. Return to this figure when you discuss tobacco and radiation."))

    add(lambda i, t: make_ipp_chevrons(prs, i, t))

    add(lambda i, t: make_content(prs, "Classification by Carcinogenic Action / Behavior", [
        "By mechanism: genotoxic vs non-genotoxic",
        "By need for activation: direct-acting vs procarcinogens",
        "By role in multistage model: initiators, promoters, progressors, complete carcinogens",
        "By evidence strength: IARC Groups 1 / 2A / 2B / 3",
        "By source: occupational, environmental, dietary, pharmaceutical, lifestyle",
        "Complete carcinogen = can drive more than one stage (e.g., some PAHs in practice)",
    ], "Chemical Carcinogens", i, t,
        "Students confuse ‘classification systems.’ Put all five on the board as complementary lenses."))

    add(lambda i, t: make_image_slide(prs, "Exposure Routes — Visual",
        "exposure_routes_diagram.jpg",
        "Inhalation · Skin absorption · Ingestion — intensity × duration × susceptibility",
        "Chemical Carcinogens", i, t,
        "Point to each portal on the figure. Ask which dominates in factories (inhalation). "
        "Mention arsenic in water (ingestion) and solvent dermatitis/absorption (skin)."))

    add(lambda i, t: make_entry_cards(prs, i, t))

    add(lambda i, t: make_content(prs, "Environmental Chemical Carcinogens — Examples", [
        "Air pollution / combustion products — PAHs, fine particulates (with lung cancer evidence for pollution)",
        "Contaminated water — arsenic (skin, lung, bladder cancers)",
        "Foodborne — aflatoxin B1 (liver), nitrates/nitrosamines (context-dependent)",
        "Persistent organics — dioxins, some PCBs (long half-life, fat storage)",
        "Second-hand tobacco smoke as an environmental carcinogen in shared air",
        "Pesticide residues — population risk usually lower than high occupational exposure",
    ], "Chemical Carcinogens", i, t,
        "Distinguish high-dose occupational clusters from low-dose general population exposures."))

    add(lambda i, t: make_content(prs, "Occupational Carcinogens — High-Yield List", [
        "Asbestos → mesothelioma; also lung and several other sites",
        "Benzene → leukemia",
        "Vinyl chloride → hepatic angiosarcoma",
        "Aromatic amines (e.g., benzidine) → bladder cancer",
        "Wood dust → sinonasal / nasopharyngeal cancers",
        "Silica, diesel exhaust, coal tars, chromium, nickel, cadmium — site-specific risks",
        "Painting / welding work and night-shift work — discuss as recognized risk contexts",
    ], "Chemical Carcinogens", i, t,
        "Use agent → cancer flash pairs. Ask students for vinyl chloride and asbestos before revealing."))

    add(lambda i, t: make_twocol(prs, "Selected Agent–Cancer Pairs",
        "Organics",
        ["Benzene → leukemia", "PAHs → lung, skin, urinary tract",
         "Vinyl chloride → liver angiosarcoma", "Benzidine → bladder",
         "Aflatoxin B1 → hepatocellular carcinoma", "Ethanol → multiple sites (IARC 1)"],
        "Dusts / fibers / metals",
        ["Asbestos → mesothelioma + lung", "Silica → lung",
         "Wood dust → sinonasal", "Arsenic → skin, lung, bladder",
         "Chromium / nickel → lung (± nasal)", "Cadmium → lung"],
        "Chemical Carcinogens", i, t,
        "This is a high-yield exam slide. Pause and quiz."))

    add(lambda i, t: make_protection(prs, i, t))

    # TOBACCO
    add(lambda i, t: make_section(prs, "02", "Tobacco & Cancer",
                                  "Contents · mechanisms · products · MSS vs SSS", i, t,
                                  "Tobacco is the most important avoidable chemical carcinogen mixture in clinical practice."))

    add(lambda i, t: make_content(prs, "Tobacco — Scale of the Problem", [
        "Leading preventable cause of cancer death worldwide",
        "Causes ~80–90% of lung cancer deaths in many populations",
        "Also causes cancers of mouth, throat, larynx, esophagus, stomach, pancreas, kidney, bladder, cervix, colon/rectum (evidence varies by site), and AML",
        "Dose–response with pack-years",
        "Risk falls after cessation — earlier quitting yields larger benefit",
    ], "Tobacco", i, t,
        "Make it clinical: every history should include smoking status and pack-years."))

    add(lambda i, t: make_image_slide(prs, "What’s Inside Smoke? — Chemical Cocktail",
        "tobacco_chemical_cocktail.jpg",
        "~4,000 chemicals · >70 known carcinogens · nicotine sustains addiction",
        "Tobacco", i, t,
        "Narrate the bubbles: PAHs and nitrosamines damage DNA; metals and Po-210 add to the burden; "
        "nicotine is not the main mutagen — it is the addiction engine that keeps exposure chronic."))

    add(lambda i, t: make_content(prs, "Key Carcinogens in Tobacco Smoke", [
        "Polycyclic aromatic hydrocarbons (e.g., benzo[a]pyrene)",
        "Tobacco-specific nitrosamines (NNK, NNN and related)",
        "Volatile organics: benzene, 1,3-butadiene, formaldehyde, acrolein",
        "Metals: arsenic, cadmium, chromium, nickel",
        "Polonium-210 (radioactive)",
        "Tar = complex carcinogenic mixture deposited in airways",
    ], "Tobacco", i, t,
        "Link back to chemical section: many of these are classic IARC agents."))

    add(lambda i, t: make_content(prs, "Mechanism of Tobacco Carcinogenesis", [
        "Inhalation → deposition in airways and absorption of systemic toxicants",
        "Metabolic activation (CYP pathways) → DNA adducts",
        "Driver mutations accumulate (e.g., TP53, KRAS patterns historically linked to smoking)",
        "Chronic inflammation and oxidative stress promote progression",
        "Field cancerization along the aerodigestive tract",
        "Systemic delivery explains distant sites (bladder, pancreas, cervix, AML)",
    ], "Tobacco", i, t,
        "Draw the pathway on the board if time: Smoke → activation → adduct → mutation → field → cancer."))

    add(lambda i, t: make_tobacco_products(prs, i, t))

    add(lambda i, t: make_image_slide(prs, "Mainstream vs Sidestream Smoke — Visual",
        "mss_sss_smoke_diagram.jpg",
        "MSS enters the smoker; SSS rises from the tip and dominates second-hand smoke",
        "Tobacco", i, t,
        "Stand at the figure and define MSS vs SSS. Ask: which exposes children in a car? (SSS/second-hand)."))

    add(lambda i, t: make_mss_sss_cards(prs, i, t))

    add(lambda i, t: make_content(prs, "Second-Hand (Environmental) Tobacco Smoke", [
        "Mixture of SSS + exhaled MSS",
        "IARC Group 1 carcinogen",
        "Causes lung cancer in nonsmokers; also cardiovascular and respiratory disease",
        "Children: higher relative intake; linked to multiple harms",
        "No safe threshold for involuntary exposure in public-health messaging",
        "Smoke-free laws are cancer prevention, not etiquette",
    ], "Tobacco", i, t,
        "Connect to pediatrics and occupational health (hospitality workers historically heavily exposed)."))

    add(lambda i, t: make_content(prs, "Clinical Counseling Pearls — Tobacco", [
        "Ask–Advise–Assist–Arrange at every opportunity",
        "Quantify pack-years; document cessation date",
        "Synergy: tobacco + asbestos / radon / alcohol (site-dependent)",
        "Do not present ‘light’ or filtered cigarettes as safe",
        "E-cigarettes: discuss honestly — potential harm-reduction tool for some smokers, risk for youth initiation",
    ], "Tobacco", i, t,
        "Role-play one 60-second cessation advice if time allows."))

    # RADIATION
    add(lambda i, t: make_section(prs, "03", "Radiation Carcinogenesis",
                                  "Types · DNA mechanisms · environment · medical exposure", i, t,
                                  "Radiation is a physical carcinogen that creates chemical injury in DNA — bridging physics and pathology."))

    add(lambda i, t: make_radiation_types(prs, i, t))

    add(lambda i, t: make_image_slide(prs, "How Radiation Damages DNA",
        "radiation_dna_damage.jpg",
        "Direct hits vs indirect free-radical (•OH) attack — both can yield mutagenic misrepair",
        "Radiation", i, t,
        "Explain direct vs indirect action. X/γ rays act largely via water radiolysis. "
        "Double-strand breaks are critical; misrepair causes chromosomal aberrations and cancer risk."))

    add(lambda i, t: make_twocol(prs, "Radiation DNA Injury — Detail",
        "Direct action",
        ["Photon/particle energy hits DNA",
         "Base damage, crosslinks, SSB/DSB",
         "More important for high-LET radiation (α, neutrons)",
         "Clustered damage harder to repair"],
        "Indirect action",
        ["Radiolysis of water → •OH and other radicals",
         "Radicals attack DNA/proteins/lipids",
         "Dominant for low-LET X/γ rays",
         "Oxygen enhances damage (oxygen effect)"],
        "Radiation", i, t,
        "Mention that stochastic cancer risk has no sharp practical ‘safe zero’ — risk rises with dose, "
        "while tissue reactions (deterministic) have thresholds."))

    add(lambda i, t: make_content(prs, "From Misrepair to Cancer", [
        "Misrepaired DSB → chromosomal aberrations / mutations",
        "Hits to oncogenes, tumor suppressors, or genome-maintenance genes",
        "Genomic instability may persist in progeny cells",
        "Latency: leukemias earlier; solid tumors often decades later",
        "Human evidence: atomic-bomb survivors, medical cohorts, radon miners, UV epidemiology",
    ], "Radiation", i, t,
        "Age at exposure matters — children often have higher lifetime risk for many tissues."))

    add(lambda i, t: make_content(prs, "Tissue Sensitivity Highlights", [
        "Highly sensitive: bone marrow (leukemia), thyroid (especially children), breast, lung",
        "UV targets skin → BCC, SCC, melanoma",
        "Radon targets lung — α-particle dose to bronchi",
        "Fractionation and dose rate modify biologic effect",
        "Combined exposures (e.g., smoking + radon) can be synergistic",
    ], "Radiation", i, t,
        "Ask: why is thyroid cancer a classic after childhood neck irradiation / fallout?"))

    add(lambda i, t: make_env_med_radiation(prs, i, t))

    add(lambda i, t: make_content(prs, "Medical Radiation — What Future Doctors Must Do", [
        "Justify every exam: will the result change management?",
        "Optimize: lowest dose compatible with diagnostic quality (ALARA)",
        "Prefer non-ionizing alternatives when equally appropriate (US/MRI)",
        "Special caution: pregnancy, children, high-dose fluoroscopy",
        "Radiotherapy: curative/palliative benefit outweighs secondary-cancer risk when indicated",
        "Communicate risks honestly without scaring patients away from needed care",
    ], "Radiation", i, t,
        "This is a professionalism slide as much as a science slide."))

    add(lambda i, t: make_content(prs, "Chemicals + Tobacco + Radiation — Shared Biology", [
        "All can initiate DNA damage (adducts, strand breaks, oxidative lesions)",
        "Promotion via inflammation and proliferation is a shared chapter",
        "Real patients stack exposures: smoker + radon + occupational dust",
        "Prevention packages beat single slogans",
        "Your clinical habit: exposure history + counseling + appropriate imaging",
    ], "Integration", i, t,
        "Draw a simple Venn on the board: overlapping mechanisms, distinct sources."))

    # CLOSE
    add(lambda i, t: make_section(prs, "04", "Prevention & Close",
                                  "Turn mechanisms into action", i, t,
                                  "Finish with agency — most of these risks are modifiable."))

    add(lambda i, t: make_content(prs, "Integrated Prevention Checklist", [
        "Never start / stop tobacco; eliminate second-hand smoke exposure",
        "Workplace: substitution, ventilation, training, PPE, regulation",
        "Sun protection; avoid tanning beds",
        "Test/mitigate radon where relevant",
        "Limit alcohol; healthy weight and diet patterns",
        "Vaccination where relevant (HPV, HBV) — infection-related cancer prevention",
        "Rational medical imaging — justification + ALARA",
    ], "Prevention", i, t,
        "Have the class shout items before you click through."))

    add(lambda i, t: make_takeaways(prs, i, t))

    add(lambda i, t: make_content(prs, "Selected References", [
        "IARC Monographs — https://monographs.iarc.who.int/",
        "WHO / IARC World Cancer Reports; tobacco & radiation fact sheets",
        "Robbins & Kumar — Basic Pathology (neoplasia / carcinogenesis)",
        "DeVita — Cancer: Principles & Practice of Oncology",
        "Perez & Brady / Gunderson — Radiation Oncology principles texts",
        "Doll R. avoidable causes of cancer (classic teaching estimates)",
        "National cancer institute / NIEHS environmental carcinogen resources",
    ], "References", i, t,
        "Point students to IARC as the living classification source."))

    total = len(builders)
    global LECTURE_NOTES
    LECTURE_NOTES = []
    for idx, fn in enumerate(builders, 1):
        fn(idx, total)

    path = OUT / "Year3_Carcinogenesis_Chemicals_Tobacco_Radiation.pptx"
    prs.save(path)
    print(f"Saved PPTX: {path.name} ({total} slides, {len(LECTURE_NOTES)} notes)")
    return path, total


def build_docx(total_slides):
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = DocPt(11)

    def H(text, level=1):
        h = doc.add_heading(text, level=level)
        for r in h.runs:
            r.font.color.rgb = DocRGB(0x14, 0x32, 0x42)

    def P(text, bold=False, italic=False):
        p = doc.add_paragraph()
        r = p.add_run(text)
        r.bold = bold
        r.italic = italic

    H("Carcinogenesis: Chemicals, Tobacco & Radiation", 0)
    P("Year 3 Medical Student Lecture — Full Speaking Script & Slide Notes", italic=True)
    P("Prof. Dr. Ş. Bilge Gürsel")
    P("Ondokuz Mayıs University Faculty of Medicine · Department of Radiation Oncology")
    P(f"Companion PowerPoint: Year3_Carcinogenesis_Chemicals_Tobacco_Radiation.pptx ({total_slides} slides)")
    P("")

    H("How to Use This Document", 1)
    P("Each section below matches one PowerPoint slide. The text is what you can say aloud "
      "(also embedded as speaker notes under every slide). Visual slides include flowcharts and diagrams — "
      "narrate the figure rather than reading bullets.")
    P("")

    H("Suggested Timing (≈ 70–90 minutes)", 1)
    items = [
        "0–8 min — Introduction & objectives",
        "8–40 min — Chemical carcinogens (mechanisms, IARC, exposures, occupation, protection)",
        "40–65 min — Tobacco (contents, mechanisms, products, MSS/SSS, counseling)",
        "65–85 min — Radiation (types, DNA injury, environment, medical exposure)",
        "85–90 min — Integrated prevention & take-home messages",
    ]
    for it in items:
        doc.add_paragraph(it, style="List Bullet")

    H("Slide-by-Slide Speaking Notes", 1)
    for num, title, notes in LECTURE_NOTES:
        H(f"Slide {num}. {title}", 2)
        P(notes)
        P("")

    H("Optional Mini-Cases for Class", 1)
    cases = [
        "A plastics worker develops hepatic angiosarcoma — which classic agent?",
        "Why can colorectal cancer rates change within one generation after migration?",
        "Parents smoke in a car with children — explain MSS vs SSS and second-hand risk.",
        "A colleague wants a CT ‘just in case’ with no clinical question — apply justification/ALARA.",
        "Name one genotoxic and one non-genotoxic carcinogenic mechanism.",
        "Is nicotine the main DNA-damaging carcinogen in cigarettes? Why does it still matter?",
    ]
    for c in cases:
        doc.add_paragraph(c, style="List Bullet")

    H("Key References", 1)
    refs = [
        "IARC Monographs — https://monographs.iarc.who.int/",
        "WHO / IARC resources on tobacco and radiation",
        "Robbins & Kumar — Basic Pathology",
        "DeVita — Cancer: Principles & Practice of Oncology",
        "Perez & Brady — Principles and Practice of Radiation Oncology",
        "Doll R. classic analyses of avoidable cancer causes",
        "NCI / NIEHS environmental carcinogen resources",
    ]
    for r in refs:
        doc.add_paragraph(r, style="List Bullet")

    path = OUT / "Year3_Carcinogenesis_Lecture_Script.docx"
    doc.save(path)
    print(f"Saved DOCX: {path.name}")
    return path


def deploy(pptx, docx):
    DESKTOP_EDU.mkdir(parents=True, exist_ok=True)
    ARTIFACTS.mkdir(parents=True, exist_ok=True)
    targets = [
        DESKTOP_EDU / "Year3 Carcinogenesis - Chemicals Tobacco Radiation.pptx",
        DESKTOP_EDU / "Year3 Carcinogenesis - Lecture Script.docx",
        ARTIFACTS / "Year3_Carcinogenesis_Chemicals_Tobacco_Radiation.pptx",
        ARTIFACTS / "Year3_Carcinogenesis_Lecture_Script.docx",
    ]
    for dest in targets:
        src = pptx if dest.suffix.lower() == ".pptx" else docx
        shutil.copy2(src, dest)
        print("→", dest)
    print("\nDesktop/Eğitim contents:")
    for f in sorted(DESKTOP_EDU.iterdir()):
        print(f"  {f.name} ({f.stat().st_size} bytes)")


if __name__ == "__main__":
    pptx, n = build_pptx()
    docx = build_docx(n)
    deploy(pptx, docx)
    print(f"\nDone: {n} slides with speaker notes + Word script.")
